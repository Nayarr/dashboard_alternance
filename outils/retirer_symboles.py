"""Remplace les symboles decoratifs du dossier par du texte.

Les caracteres ✉ ☎ ⌂ ne sont pas couverts par Calibri : Word bascule sur Segoe
UI Emoji et embarque la police entiere dans le PDF, faute de savoir la
sous-ensembler. Cout mesure : 3,8 Mo pour trois caracteres.

    python outils/retirer_symboles.py
"""

import shutil
from pathlib import Path

from docx import Document

BASE = Path(__file__).resolve().parent.parent
SOURCE = BASE / "templates" / "dossier" / "dossier_base.docx"

REMPLACEMENTS = {
    "✉": "Email :",        # ✉
    "☎": "Téléphone :",    # ☎
    "⌂": "Adresse :",      # ⌂
    "☏": "Téléphone :",    # ☏
    "\U0001F4E7": "Email :",
    "\U0001F4F1": "Téléphone :",
    "\U0001F3E0": "Adresse :",
}


def symboles_presents(doc):
    trouves = {}
    for para in doc.paragraphs:
        for c in para.text:
            if ord(c) > 0x2000 and c not in "‘’“”–—…· ":
                trouves.setdefault(c, 0)
                trouves[c] += 1
    return trouves


def main():
    backup = SOURCE.with_suffix(".docx.avant_symboles")
    if not backup.exists():
        shutil.copy2(SOURCE, backup)

    doc = Document(SOURCE)

    presents = symboles_presents(doc)
    if presents:
        print("Symboles hors Calibri detectes :")
        for c, n in presents.items():
            print(f"   U+{ord(c):04X} {c!r} x{n}")

    modifs = 0
    for para in doc.paragraphs:
        for run in para.runs:
            texte = run.text
            for symbole, libelle in REMPLACEMENTS.items():
                if symbole in texte:
                    texte = texte.replace(symbole, libelle)
                    modifs += 1
            if texte != run.text:
                run.text = texte
                # La police emoji reste parfois collee au run : on force Calibri
                run.font.name = "Calibri"

    doc.save(SOURCE)
    print(f"\n{modifs} remplacement(s) — {SOURCE.name} mis a jour")

    restants = symboles_presents(Document(SOURCE))
    if restants:
        print("Symboles restants (a traiter si le PDF reste lourd) :")
        for c, n in restants.items():
            print(f"   U+{ord(c):04X} {c!r} x{n}")


if __name__ == "__main__":
    main()
