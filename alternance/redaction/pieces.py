"""Fabrique les pieces jointes d'une candidature : lettre DOCX -> PDF, dossier PDF.

La conversion passe par Word en COM (Word 16 present sur le poste), qui preserve
la mise en forme exactement, contrairement a un rendu maison.

    python cli.py pieces --offre 925
"""

import argparse
import shutil
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from alternance import chemins
from alternance import config
from alternance import db

BASE = config.RACINE
LETTRES = BASE / "lettres"
DOSSIER_DOCX = BASE / "templates" / "dossier" / "dossier_base.docx"
CV_PDF = chemins.CV

ACCENT = RGBColor(0x58, 0x1C, 0x87)   # meme violet que le CV LaTeX
GRIS = RGBColor(0x4B, 0x55, 0x63)


def slug(texte):
    return "".join(c if c.isalnum() else "_" for c in (texte or "").lower()).strip("_")


def construire_lettre_docx(corps, entreprise, destination):
    """Compose la lettre en DOCX : en-tete de coordonnees puis corps."""
    d = docx.Document()

    section = d.sections[0]
    section.top_margin = section.bottom_margin = Pt(50)
    section.left_margin = section.right_margin = Pt(60)

    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    # En-tete : identite
    p = d.add_paragraph()
    r = p.add_run(config.PROFIL["nom"])
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = ACCENT
    p.paragraph_format.space_after = Pt(2)

    p = d.add_paragraph()
    r = p.add_run("Étudiant en BUT Informatique  ·  Recherche d'alternance")
    r.font.size = Pt(10)
    r.font.color.rgb = GRIS
    p.paragraph_format.space_after = Pt(2)

    p = d.add_paragraph()
    # Le code postal est facultatif : sans lui, on evite "Paris ()" dans l'en-tete
    cp = config.PROFIL.get("code_postal") or ""
    localisation = f"{config.PROFIL['ville']} ({cp})" if cp else config.PROFIL["ville"]
    r = p.add_run(f"{localisation}  ·  "
                  f"{config.PROFIL['telephone_affiche']}  ·  {config.PROFIL['email']}")
    r.font.size = Pt(9.5)
    r.font.color.rgb = GRIS
    p.paragraph_format.space_after = Pt(18)

    # Destinataire
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(f"À l'attention du service recrutement\n{entreprise}")
    r.font.size = Pt(10)
    p.paragraph_format.space_after = Pt(20)

    # Corps
    for bloc in [b.strip() for b in corps.split("\n\n") if b.strip()]:
        p = d.add_paragraph(bloc)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(9)
        p.paragraph_format.line_spacing = 1.12

    d.save(destination)
    return destination


def docx_vers_pdf(source, destination, pour_ecran=False):
    """Conversion via Word COM, en qualite impression.

    OptimizeFor=1 (ecran) reechantillonne les images a 96 ppp : le CV insere en
    page 4 du dossier tombait de 1240x1755 a 619x877, visiblement degrade. On
    reste donc en qualite impression (0).

    Le poids du dossier ne venait de toute facon pas des images mais de Segoe UI
    Emoji embarquee en entier, corrige par outils/retirer_symboles.py.
    """
    import win32com.client

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        doc = word.Documents.Open(str(Path(source).resolve()))
        doc.ExportAsFixedFormat(
            OutputFileName=str(Path(destination).resolve()),
            ExportFormat=17,               # wdExportFormatPDF
            OpenAfterExport=False,
            OptimizeFor=1 if pour_ecran else 0,   # 1 = wdExportOptimizeForOnScreen
            CreateBookmarks=0,
            DocStructureTags=True,
        )
        doc.Close(False)
    finally:
        word.Quit()
    return destination


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--offre", type=int, required=True)
    args = p.parse_args()

    conn = db.connect()
    o = conn.execute("SELECT * FROM offres WHERE id = ?", (args.offre,)).fetchone()
    if o is None:
        sys.exit(f"offre {args.offre} introuvable")

    dossier = chemins.dossier_candidature(o, creer=True)

    texte = chemins.lettre_txt(o)
    if not texte.exists():
        sys.exit(f"lettre absente : {texte}\nRediger la lettre avant de generer.")
    corps = texte.read_text(encoding="utf-8")

    print(f"Offre #{o['id']} — {o['entreprise']}")

    lettre_docx = construire_lettre_docx(corps, o["entreprise"], dossier / "lettre.docx")
    print(f"  lettre DOCX  {lettre_docx.name}")

    lettre_pdf = docx_vers_pdf(lettre_docx, dossier / "lettre.pdf")
    print(f"  lettre PDF   {Path(lettre_pdf).name}")

    dossier_pdf = dossier / "dossier_de_motivation.pdf"
    docx_vers_pdf(DOSSIER_DOCX, dossier_pdf)
    print(f"  dossier PDF  {dossier_pdf.name}")

    if CV_PDF.exists():
        shutil.copy2(CV_PDF, dossier / CV_PDF.name)
        print("  CV PDF       %s" % CV_PDF.name)
    else:
        print(f"  CV ABSENT    deposer le PDF dans {CV_PDF}")

    conn.close()
    print(f"\nPieces pretes dans {dossier}")


if __name__ == "__main__":
    main()
